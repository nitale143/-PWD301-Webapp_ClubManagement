"""Luồng nghiệp vụ chính qua giao diện web: quỹ theo kỳ, sự kiện – task – tính
hoạt động, danh sách/xuất Excel thành viên và các lệnh CLI quản trị."""
import io
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from flask import g
from openpyxl import load_workbook

from tests.support import csrf_for
from app import create_app, db
from app.models import (
    Ban, Event, FundDue, FundPeriod, FundTransaction, MemberActivity, MemberRecord, Role, Task, User,
    UserBan,
)
from app.services.common import club_now

PASSWORD = "Matkhau-2026"


class CoreFlowTests(unittest.TestCase):
    USERS = ((1, "CN"), (2, "TB"), (3, "THANH_VIEN"), (4, "THUKY_THUQUY"))

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.app = create_app({
            "TESTING": True, "SECRET_KEY": "core-flow-secret-key-0123456789abcdef",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{Path(self.directory.name) / 'core.db'}",
            "AI_PROVIDER": "rules", "SCHEDULER_ENABLED": False})
        self.client = self.app.test_client()
        self.context = self.app.app_context()
        self.context.push()
        db.session.add(Ban(id=1, ten_ban="Kỹ thuật"))
        for user_id, role in self.USERS:
            user = User(id=user_id, mssv=f"SV{user_id:04d}", ho_ten=f"User {user_id}",
                        ngay_sinh=date(2003, 1, 1), sdt="0912345678",
                        email=f"user{user_id}@example.com", status="approved", chuc_vu=role)
            user.password = PASSWORD
            db.session.add(user)
            db.session.add(MemberRecord(user_id=user_id, ngay_gia_nhap=date(2024, 1, 1)))
            db.session.add(UserBan(user_id=user_id, ban_id=1))
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.engine.dispose()
        self.context.pop()
        self.directory.cleanup()

    def login(self, user_id):
        with self.client.session_transaction() as session:
            session["_user_id"] = str(user_id)
            session["_fresh"] = True
        g.pop("_login_user", None)

    def post(self, path, data=None):
        g.pop("_login_user", None)
        response = self.client.post(path, data={"csrf_token": csrf_for(self.client), **(data or {})})
        g.pop("_login_user", None)
        return response

    def get(self, path):
        g.pop("_login_user", None)
        return self.client.get(path)

    # ---------------- Quỹ theo kỳ ----------------
    def test_fund_period_ledger_and_dues_follow_permissions(self):
        today = club_now().date()
        self.login(2)   # Trưởng ban: không có quyền quỹ
        self.assertEqual(self.post("/funds/period/new", {"ten_ky": "Fall", "ngay_bat_dau": today.isoformat()}).status_code, 403)

        self.login(4)   # Thư ký/Thủ quỹ
        self.post("/funds/period/new", {"ten_ky": "", "ngay_bat_dau": today.isoformat()})
        self.assertIsNone(FundPeriod.query.first())
        self.assertEqual(self.post("/funds/period/new", {
            "ten_ky": "Fall 2026", "ngay_bat_dau": today.isoformat(),
            "ngay_ket_thuc": (today + timedelta(days=90)).isoformat()}).status_code, 302)
        period = FundPeriod.query.filter_by(is_current=True).one()
        self.assertEqual({due.user_id for due in FundDue.query.filter_by(period_id=period.id)}, {1, 2, 3, 4})

        self.post("/funds/transaction/new", {"danh_muc": "x" * 101, "ngay": today.isoformat(), "so_tien": "1000"})
        self.post("/funds/transaction/new", {"danh_muc": "Thu quỹ", "ngay": today.isoformat(), "so_tien": "abc"})
        self.assertEqual(FundTransaction.query.count(), 0)
        self.post("/funds/transaction/new", {"danh_muc": "Thu quỹ", "ngay": today.isoformat(), "so_tien": "500000"})
        self.post("/funds/transaction/new", {"danh_muc": "In ấn", "noi_dung": "Poster", "ngay": today.isoformat(), "so_tien": "-120000"})
        self.assertEqual(sorted(int(tx.so_tien) for tx in FundTransaction.query), [-120000, 500000])

        self.assertEqual(self.post("/funds/due/3/toggle").status_code, 302)
        self.assertTrue(FundDue.query.filter_by(period_id=period.id, user_id=3).one().da_dong)
        with patch("app.routes.funds.send_email", return_value=True) as mail:
            self.post("/funds/nhac-dong-quy")
        self.assertEqual(mail.call_count, 3)

        workbook = load_workbook(io.BytesIO(self.get("/funds/export").data))
        self.assertEqual(workbook["Thu chi"].max_row, 3)
        self.assertEqual(workbook["Đóng quỹ"].max_row, 5)

        self.login(3)   # Thành viên chỉ thấy dòng đóng quỹ của mình, không thấy sổ thu chi
        workbook = load_workbook(io.BytesIO(self.get("/funds/export").data))
        self.assertEqual(workbook["Thu chi"].max_row, 1)
        self.assertEqual([row[1] for row in workbook["Đóng quỹ"].iter_rows(min_row=2, values_only=True)], ["SV0003"])
        self.assertEqual(self.post("/funds/due/3/toggle").status_code, 403)

    # ---------------- Sự kiện → task → tính hoạt động ----------------
    def test_event_task_review_and_activity_calculation(self):
        now = club_now()
        db.session.add(FundPeriod(id=1, ten_ky="Kỳ hiện tại", ngay_bat_dau=(now - timedelta(days=30)).date(),
                                  ngay_ket_thuc=(now + timedelta(days=60)).date(), is_current=True))
        db.session.commit()
        self.login(4)   # Thư ký/Thủ quỹ có quyền sự kiện & task
        starts, ends = now + timedelta(days=1), now + timedelta(days=2)
        bad = self.post("/event/new", {"ten_su_kien": "Workshop", "ma_su_kien": "bad code!",
                                       "thoi_gian_bat_dau": starts.strftime("%Y-%m-%dT%H:%M"),
                                       "thoi_gian_ket_thuc": ends.strftime("%Y-%m-%dT%H:%M")})
        self.assertIn("Mã chỉ gồm chữ".encode(), bad.data)
        created = self.post("/event/new", {"ten_su_kien": "Workshop", "ma_su_kien": "WS-01",
                                           "thoi_gian_bat_dau": starts.strftime("%Y-%m-%dT%H:%M"),
                                           "thoi_gian_ket_thuc": ends.strftime("%Y-%m-%dT%H:%M")})
        self.assertEqual(created.status_code, 302)
        event = Event.query.filter_by(ma_su_kien="WS-01").one()
        self.assertIsNone(event.detail.so_nguoi_toi_da)

        deadline = (now + timedelta(hours=30)).strftime("%Y-%m-%dT%H:%M")
        with patch("app.routes.event.send_email", return_value=True):
            self.post(f"/event/{event.id}/task/new", {"ten_task": "x" * 151, "assignee_mssv": "SV0003",
                                                      "ban_id": "1", "deadline": deadline})
            self.assertEqual(Task.query.count(), 0)
            self.post(f"/event/{event.id}/task/new", {"ten_task": "Poster", "assignee_mssv": "SV0003",
                                                      "ban_id": "1", "deadline": deadline})
            self.post(f"/event/{event.id}/task/new", {"ten_task": "Slide", "assignee_mssv": "SV0003",
                                                      "ban_id": "1", "deadline": deadline})
        poster, slide = Task.query.order_by(Task.id).all()

        self.login(3)
        with patch("app.routes.event.send_email", return_value=True):
            self.post(f"/task/{poster.id}/nop-san-pham", {"noi_dung_nop": "https://drive/poster"})
        self.assertEqual(db.session.get(Task, poster.id).trang_thai, "cho_duyet")
        self.assertEqual(self.post(f"/task/{poster.id}/cap-nhat-trang-thai", {"trang_thai": "hoan_thanh"}).status_code, 403)

        self.login(4)
        with patch("app.routes.event.send_email", return_value=True):
            self.post(f"/task/{poster.id}/cap-nhat-trang-thai", {"trang_thai": "lam_lai"})
            self.assertEqual(db.session.get(Task, poster.id).trang_thai, "cho_duyet")   # thiếu nhận xét
            self.post(f"/task/{poster.id}/cap-nhat-trang-thai", {"trang_thai": "hoan_thanh"})
            self.post(f"/task/{slide.id}/cap-nhat-trang-thai", {"trang_thai": "huy"})
        self.assertEqual(self.post(f"/event/{event.id}/ket-thuc").status_code, 302)
        activity = MemberActivity.query.filter_by(user_id=3, ban_id=1, period_id=1).one()
        self.assertEqual((activity.tong_task, activity.task_hoan_thanh, activity.phan_tram), (1, 1, 100.0))
        self.assertEqual(activity.trang_thai, "tich_cuc")

    # ---------------- Danh sách & xuất Excel thành viên ----------------
    def test_member_page_and_export_respect_member_management_permission(self):
        self.login(2)   # Trưởng ban có quyền quản lý thành viên
        page = self.get("/member?q=User&sort=mssv")
        self.assertIn(b"user3@example.com", page.data)
        sheet = load_workbook(io.BytesIO(self.get("/member/export").data)).active
        self.assertEqual(sheet.max_row, 5)
        self.assertIn("Email", [cell.value for cell in sheet[1]])
        self.assertEqual(self.get("/profile/3").status_code, 200)

        self.login(4)   # Thư ký/Thủ quỹ không quản lý thành viên → chỉ thấy bản thân
        page = self.get("/member")
        self.assertNotIn(b"user3@example.com", page.data)
        sheet = load_workbook(io.BytesIO(self.get("/member/export").data)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertNotIn("Email", [cell.value for cell in sheet[1]])
        self.assertEqual(self.get("/profile/3").status_code, 403)


class CliTests(unittest.TestCase):
    def test_init_db_create_cn_and_sync_roles(self):
        with tempfile.TemporaryDirectory() as directory:
            app = create_app({"TESTING": True, "AUTO_CREATE_TABLES": False, "SCHEDULER_ENABLED": False,
                              "SECRET_KEY": "cli-test-secret-key-0123456789abcdef",
                              "SQLALCHEMY_DATABASE_URI": f"sqlite:///{directory}/cli.db"})
            runner = app.test_cli_runner()
            migrations = str(Path(__file__).resolve().parents[1] / "migrations")
            with app.app_context():
                app.extensions["migrate"].directory = migrations
                result = runner.invoke(args=["init-db"])
                self.assertEqual(result.exit_code, 0, result.output)
                self.assertEqual(Ban.query.count(), 3)

                weak = runner.invoke(args=["create-cn"], input="CN00001\nChu Nhiem\ncn@example.com\nabc\nabc\n")
                self.assertNotEqual(weak.exit_code, 0)
                created = runner.invoke(args=["create-cn"],
                                        input=f"CN00001\nChu Nhiem\nCN@Example.com\n{PASSWORD}\n{PASSWORD}\n")
                self.assertEqual(created.exit_code, 0, created.output)
                user = User.query.filter_by(mssv="CN00001").one()
                self.assertEqual((user.chuc_vu, user.status, user.email), ("CN", "approved", "cn@example.com"))

                role = Role.query.filter_by(name="TB").one()
                role.permissions = 0
                db.session.commit()
                synced = runner.invoke(args=["sync-roles"])
                self.assertEqual(synced.exit_code, 0, synced.output)
                db.session.refresh(role)
                self.assertNotEqual(role.permissions, 0)
                db.session.remove()
                db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
