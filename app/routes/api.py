"""JSON endpoints for the expanded club workflows; the existing pages stay intact."""

import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from io import StringIO
import os
import secrets

from flask import Blueprint, current_app, jsonify, request, send_file
from flask_login import current_user
from openpyxl import Workbook, load_workbook
import qrcode
from PIL import Image, UnidentifiedImageError
from sqlalchemy import or_

from app import db
from app.models import (
    ActivityPoint, AIProposal, Attendance, AuditLog, Ban, ClubSetting, Event, EventDetail,
    EventRegistration, EventTargetBan, FundCollection, FundPayment,
    FundTargetBan, MemberRecord, PaymentEvidence, User, UserBan,
)
from app.services.access import (
    approved, can_manage_member, can_view_member, can_view_all_members,
    is_admin, is_board, is_treasurer, require, visible_members,
)
from app.services.activities import activity_stats, manual_point, DEFAULT_POINTS, DEFAULT_THRESHOLDS
from app.services.assistant import get_provider
from app.services.proposals import create_proposal, proposal_preview, review_proposal
from app.services.common import DomainError, audit, member_record, parse_datetime
from app.services.common import parse_date
from app.services.events import (
    cancel_registration, feedback, manual_attendance, qr_checkin, qr_token, register,
)
from app.services.funds import (
    adjust_fund, balance, create_collection, record_payment, review_payment,
    target_members,
)
from app.validators import EMAIL_RE, MSSV_RE, PHONE_RE


api_bp = Blueprint("api", __name__, url_prefix="/api/v1")


@api_bp.before_request
def authenticate():
    if not current_user.is_authenticated:
        return jsonify(error="Vui lòng đăng nhập."), 401
    if not approved(current_user):
        return jsonify(error="Tài khoản chưa được duyệt hoặc đã lưu trữ."), 403


@api_bp.errorhandler(DomainError)
def domain_error(error):
    db.session.rollback()
    return jsonify(error=str(error)), error.status


@api_bp.errorhandler(404)
def not_found(_error):
    return jsonify(error="Không tìm thấy dữ liệu."), 404


def body():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise DomainError("Dữ liệu JSON không hợp lệ.")
    return data


def paginate(query):
    try:
        page = int(request.args.get("page", 1))
        per_page = int(request.args.get("page_size", 20))
    except ValueError:
        raise DomainError("Trang hoặc số bản ghi mỗi trang không hợp lệ.") from None
    if page < 1 or not 1 <= per_page <= 100:
        raise DomainError("Trang hoặc số bản ghi mỗi trang không hợp lệ.")
    return query.paginate(page=page, per_page=per_page, error_out=False)


def listing(page, serializer):
    return jsonify(items=[serializer(row) for row in page.items], page=page.page,
                   page_size=page.per_page, total=page.total, pages=page.pages)


def member_json(user):
    record = member_record(user)
    return {"id": user.id, "code": user.mssv, "name": user.ho_ten,
            "email": user.email, "phone": user.sdt, "avatar": user.avatar_url,
            "birth_date": user.ngay_sinh.isoformat(), "role": user.chuc_vu,
            "departments": [link.ban_id for link in user.ban_links],
            "joined_at": record.ngay_gia_nhap.isoformat(),
            "status": record.tinh_trang, "note": record.ghi_chu}


def _create_member(data):
    require(is_board(current_user))
    code = str(data.get("code", "")).strip()
    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip()
    phone = str(data.get("phone", "")).strip()
    password = str(data.get("password", ""))
    if not MSSV_RE.fullmatch(code) or User.query.filter_by(mssv=code).first():
        raise DomainError("Mã thành viên không hợp lệ hoặc đã tồn tại.")
    if len(name) < 2 or len(name) > 100 or not EMAIL_RE.fullmatch(email) or not PHONE_RE.fullmatch(phone):
        raise DomainError("Họ tên, email hoặc số điện thoại không hợp lệ.")
    if len(password) < 8:
        raise DomainError("Mật khẩu ban đầu phải có ít nhất 8 ký tự.")
    birth = parse_date(data.get("birth_date"), "Ngày sinh")
    joined = parse_date(data.get("joined_at"), "Ngày gia nhập") if data.get("joined_at") else datetime.utcnow().date()
    if birth >= datetime.utcnow().date() or joined < birth:
        raise DomainError("Ngày sinh hoặc ngày gia nhập không hợp lệ.")
    raw_departments = data.get("department_ids", [])
    if isinstance(raw_departments, str):
        raw_departments = raw_departments.split(",")
    try:
        department_ids = list({int(value) for value in raw_departments})
    except (TypeError, ValueError):
        raise DomainError("Danh sách ban không hợp lệ.") from None
    if not department_ids or Ban.query.filter(Ban.id.in_(department_ids)).count() != len(department_ids):
        raise DomainError("Danh sách ban không hợp lệ.")
    if current_user.chuc_vu == "TB" and not set(department_ids).issubset(current_user.quan_ly_ban_ids()):
        raise DomainError("Trưởng ban chỉ được thêm thành viên vào ban mình quản lý.", 403)
    user = User(mssv=code, ho_ten=name, ngay_sinh=birth, sdt=phone, email=email,
                status="approved", chuc_vu="THANH_VIEN")
    user.set_password(password)
    db.session.add(user)
    db.session.flush()
    db.session.add(MemberRecord(user_id=user.id, ngay_gia_nhap=joined,
                                tinh_trang=str(data.get("status", "active")),
                                ghi_chu=str(data.get("note", ""))))
    for department_id in department_ids:
        db.session.add(UserBan(user_id=user.id, ban_id=department_id))
    audit(current_user, "member_create", user)
    return user


def fund_json(fund, member=None):
    item = {"id": fund.id, "name": fund.ten_khoan_thu, "type": fund.loai,
            "amount": str(fund.so_tien), "start_date": fund.ngay_bat_dau.isoformat(),
            "due_date": fund.han_dong.isoformat(), "status": fund.trang_thai,
            "description": fund.mo_ta, "applies_to_all": fund.tat_ca_thanh_vien}
    if member:
        try:
            values = balance(fund, member)
            item["balance"] = {key: str(value) if isinstance(value, Decimal) else value
                               for key, value in values.items()}
        except DomainError:
            pass
    return item


def payment_json(payment):
    return {"id": payment.id, "fund_id": payment.fund_id, "member_id": payment.user_id,
            "amount": str(payment.so_tien), "paid_at": payment.ngay_dong.isoformat(),
            "method": payment.phuong_thuc, "transaction_code": payment.ma_giao_dich,
            "status": payment.trang_thai, "reviewed_by": payment.xac_nhan_boi_id,
            "note": payment.ghi_chu,
            "evidence_ids": [row.id for row in PaymentEvidence.query.filter_by(payment_id=payment.id)]}


def event_json(event):
    detail = db.session.get(EventDetail, event.id)
    return {"id": event.id, "code": event.ma_su_kien, "name": event.ten_su_kien,
            "start_at": event.thoi_gian_bat_dau.isoformat(),
            "end_at": event.thoi_gian_ket_thuc.isoformat(), "status": event.trang_thai,
            "type": detail.loai if detail else "other",
            "content": detail.noi_dung if detail else "",
            "location": detail.dia_diem if detail else "",
            "image": detail.anh_dai_dien if detail else "",
            "registration_deadline": detail.han_dang_ky.isoformat() if detail and detail.han_dang_ky else None,
            "capacity": detail.so_nguoi_toi_da if detail else None,
            "manager_id": detail.phu_trach_id if detail else None,
            "fund_id": detail.fund_id if detail else None}


def can_manage_event(event):
    return is_board(current_user) and (
        current_user.chuc_vu in {"CN", "PCN"} or event.tao_boi_id == current_user.id
    )


@api_bp.get("/members")
def members():
    query = visible_members(current_user)
    search = request.args.get("q", "").strip()
    if search:
        query = query.filter(or_(User.ho_ten.ilike(f"%{search}%"), User.mssv.ilike(f"%{search}%")))
    if request.args.get("department_id"):
        try:
            department_id = int(request.args["department_id"])
        except ValueError:
            raise DomainError("Mã ban không hợp lệ.") from None
        query = query.join(UserBan).filter(UserBan.ban_id == department_id)
    if request.args.get("status"):
        query = query.join(MemberRecord).filter(MemberRecord.tinh_trang == request.args["status"])
    return listing(paginate(query.order_by(User.id)), member_json)


@api_bp.post("/members")
def member_create():
    data = body()
    if data.get("status", "active") not in {"active", "paused", "left"}:
        raise DomainError("Trạng thái thành viên không hợp lệ.")
    user = _create_member(data)
    db.session.commit()
    return jsonify(member_json(user)), 201


@api_bp.post("/members/import")
def members_import():
    require(is_board(current_user))
    upload = request.files.get("file")
    if not upload:
        raise DomainError("Vui lòng chọn tệp CSV hoặc XLSX.")
    content = upload.read(3 * 1024 * 1024 + 1)
    if len(content) > 3 * 1024 * 1024:
        raise DomainError("Tệp nhập vượt 3 MB.")
    name = (upload.filename or "").lower()
    if name.endswith(".csv"):
        try:
            rows = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
        except (UnicodeDecodeError, csv.Error):
            raise DomainError("Tệp CSV phải được mã hóa UTF-8.") from None
    elif name.endswith(".xlsx"):
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(cell or "").strip() for cell in next(values)]
            rows = [dict(zip(headers, row)) for row in values]
            workbook.close()
        except (ValueError, OSError, StopIteration):
            raise DomainError("Tệp XLSX không hợp lệ.") from None
    else:
        raise DomainError("Chỉ hỗ trợ tệp CSV hoặc XLSX.")
    if not 1 <= len(rows) <= 500:
        raise DomainError("Tệp nhập phải có từ 1 đến 500 thành viên.")
    required = {"code", "name", "birth_date", "email", "phone", "department_ids", "password"}
    if not required.issubset(rows[0]):
        raise DomainError("Thiếu cột bắt buộc: " + ", ".join(sorted(required - set(rows[0]))))
    ids = []
    for number, row in enumerate(rows, start=2):
        try:
            if row.get("status", "active") not in {"active", "paused", "left"}:
                raise DomainError("Trạng thái thành viên không hợp lệ.")
            user = _create_member(row)
            ids.append(user.id)
        except DomainError as error:
            raise DomainError(f"Dòng {number}: {error}", error.status) from None
    audit(current_user, "members_import", current_user, {"count": len(ids)})
    db.session.commit()
    return jsonify(imported=len(ids), member_ids=ids), 201


@api_bp.get("/members/<int:user_id>")
def member_detail(user_id):
    user = db.get_or_404(User, user_id)
    require(can_view_member(current_user, user))
    return jsonify(member_json(user))


@api_bp.patch("/members/<int:user_id>")
def member_update(user_id):
    user = db.get_or_404(User, user_id)
    require(can_manage_member(current_user, user))
    data = body()
    record = member_record(user)
    if "name" in data:
        name = str(data["name"]).strip()
        if not name or len(name) > 100:
            raise DomainError("Họ tên không hợp lệ.")
        user.ho_ten = name
    if "phone" in data:
        phone = str(data["phone"]).strip()
        if not phone or len(phone) > 15:
            raise DomainError("Số điện thoại không hợp lệ.")
        user.sdt = phone
    if "status" in data:
        if data["status"] not in {"active", "paused", "left"}:
            raise DomainError("Trạng thái thành viên không hợp lệ.")
        record.tinh_trang = data["status"]
    if "note" in data:
        record.ghi_chu = str(data["note"] or "").strip()
    audit(current_user, "member_update", record, {"member_id": user.id})
    db.session.commit()
    return jsonify(member_json(user))


@api_bp.post("/members/<int:user_id>/archive")
def member_archive(user_id):
    user = db.get_or_404(User, user_id)
    require(is_admin(current_user) and user.id != current_user.id)
    record = member_record(user)
    if record.luu_tru_luc:
        raise DomainError("Thành viên đã được lưu trữ.")
    record.tinh_trang = "left"
    record.luu_tru_luc = datetime.utcnow()
    audit(current_user, "member_archive", record, {"member_id": user.id})
    db.session.commit()
    return jsonify(message="Đã lưu trữ thành viên.")


@api_bp.get("/members/<int:user_id>/activity")
def member_activity(user_id):
    user = db.get_or_404(User, user_id)
    require(can_view_member(current_user, user))
    result = activity_stats(user)
    result["latest"] = result["latest"].isoformat() if result["latest"] else None
    return jsonify(result)


@api_bp.post("/members/<int:user_id>/points")
def member_points(user_id):
    user = db.get_or_404(User, user_id)
    require(is_board(current_user) and can_manage_member(current_user, user))
    data = body()
    event = db.get_or_404(Event, data["event_id"]) if data.get("event_id") else None
    point = manual_point(current_user, user, data.get("points"), data.get("reason"), event)
    return jsonify(id=point.id, points=point.so_diem), 201


@api_bp.get("/funds")
def funds():
    query = FundCollection.query.order_by(FundCollection.id.desc())
    if not is_treasurer(current_user):
        fund_ids = [fund.id for fund in query if target_members(fund).filter(User.id == current_user.id).first()]
        query = FundCollection.query.filter(FundCollection.id.in_(fund_ids)).order_by(FundCollection.id.desc())
    return listing(paginate(query), lambda fund: fund_json(fund, current_user))


@api_bp.post("/funds")
def fund_create():
    fund = create_collection(current_user, body())
    return jsonify(fund_json(fund)), 201


@api_bp.get("/funds/<int:fund_id>")
def fund_detail(fund_id):
    fund = db.get_or_404(FundCollection, fund_id)
    require(is_treasurer(current_user) or target_members(fund).filter(User.id == current_user.id).first() is not None)
    return jsonify(fund_json(fund, current_user))


@api_bp.patch("/funds/<int:fund_id>/status")
def fund_status(fund_id):
    require(is_treasurer(current_user))
    fund = db.get_or_404(FundCollection, fund_id)
    status = body().get("status")
    allowed = {"draft": {"open", "cancelled"}, "open": {"closed", "cancelled"},
               "closed": set(), "cancelled": set()}
    if status not in allowed[fund.trang_thai]:
        raise DomainError("Không thể chuyển sang trạng thái khoản thu này.")
    fund.trang_thai = status
    audit(current_user, "fund_status", fund, {"status": status})
    db.session.commit()
    return jsonify(fund_json(fund))


@api_bp.get("/funds/<int:fund_id>/balances")
def fund_balances(fund_id):
    require(is_treasurer(current_user))
    fund = db.get_or_404(FundCollection, fund_id)
    query = target_members(fund).order_by(User.id)
    page = paginate(query)
    return listing(page, lambda user: {"member": member_json(user), "balance": fund_json(fund, user)["balance"]})


@api_bp.post("/funds/<int:fund_id>/payments")
def payment_create(fund_id):
    fund = db.get_or_404(FundCollection, fund_id)
    data = body()
    member = db.get_or_404(User, data.get("member_id", current_user.id))
    payment = record_payment(current_user, fund, member, data)
    return jsonify(payment_json(payment)), 201


@api_bp.get("/payments")
def payments():
    query = FundPayment.query
    if not is_treasurer(current_user):
        query = query.filter_by(user_id=current_user.id)
    if request.args.get("status"):
        query = query.filter_by(trang_thai=request.args["status"])
    return listing(paginate(query.order_by(FundPayment.id.desc())), payment_json)


@api_bp.post("/payments/<int:payment_id>/review")
def payment_review(payment_id):
    payment = db.get_or_404(FundPayment, payment_id)
    data = body()
    review_payment(current_user, payment, data.get("decision"), data.get("note", ""))
    return jsonify(payment_json(payment))


@api_bp.post("/payments/<int:payment_id>/evidence")
def payment_evidence_upload(payment_id):
    payment = db.get_or_404(FundPayment, payment_id)
    require(payment.user_id == current_user.id or is_treasurer(current_user))
    if payment.trang_thai != "pending":
        raise DomainError("Chỉ thêm minh chứng khi giao dịch đang chờ xác nhận.")
    upload = request.files.get("file")
    if not upload:
        raise DomainError("Vui lòng chọn ảnh minh chứng.")
    content = upload.read(3 * 1024 * 1024 + 1)
    if not content or len(content) > 3 * 1024 * 1024:
        raise DomainError("Ảnh minh chứng phải nhỏ hơn 3 MB.")
    try:
        image = Image.open(BytesIO(content))
        image.verify()
        image_format = image.format
    except (UnidentifiedImageError, OSError):
        raise DomainError("Tệp không phải ảnh hợp lệ.") from None
    suffix = {"PNG": ".png", "JPEG": ".jpg", "WEBP": ".webp"}.get(image_format)
    if not suffix:
        raise DomainError("Chỉ hỗ trợ ảnh PNG, JPEG hoặc WebP.")
    directory = current_app.config.get("PAYMENT_EVIDENCE_DIR") or os.path.join(current_app.instance_path, "payment_evidence")
    os.makedirs(directory, exist_ok=True)
    filename = secrets.token_hex(20) + suffix
    with open(os.path.join(directory, filename), "wb") as handle:
        handle.write(content)
    evidence = PaymentEvidence(payment_id=payment.id, duong_dan=filename)
    db.session.add(evidence)
    db.session.flush()
    audit(current_user, "payment_evidence_upload", evidence, {"payment_id": payment.id})
    db.session.commit()
    return jsonify(id=evidence.id), 201


@api_bp.get("/payments/evidence/<int:evidence_id>")
def payment_evidence_download(evidence_id):
    evidence = db.get_or_404(PaymentEvidence, evidence_id)
    payment = db.get_or_404(FundPayment, evidence.payment_id)
    require(payment.user_id == current_user.id or is_treasurer(current_user))
    directory = current_app.config.get("PAYMENT_EVIDENCE_DIR") or os.path.join(current_app.instance_path, "payment_evidence")
    return send_file(os.path.join(directory, evidence.duong_dan))


@api_bp.put("/funds/<int:fund_id>/adjustments/<int:user_id>")
def fund_adjustment(fund_id, user_id):
    fund = db.get_or_404(FundCollection, fund_id)
    user = db.get_or_404(User, user_id)
    adjustment = adjust_fund(current_user, fund, user, body())
    return jsonify(id=adjustment.id, type=adjustment.loai, amount=str(adjustment.so_tien))


@api_bp.get("/events")
def events():
    query = Event.query.order_by(Event.thoi_gian_bat_dau.desc())
    if request.args.get("status"):
        query = query.filter_by(trang_thai=request.args["status"])
    return listing(paginate(query), event_json)


@api_bp.post("/events")
def event_create():
    require(is_board(current_user))
    data = body()
    name = str(data.get("name", "")).strip()
    code = str(data.get("code", "")).strip()
    if not name or not code or Event.query.filter_by(ma_su_kien=code).first():
        raise DomainError("Tên hoặc mã sự kiện không hợp lệ hay đã tồn tại.")
    start = parse_datetime(data.get("start_at"), "Giờ bắt đầu")
    end = parse_datetime(data.get("end_at"), "Giờ kết thúc")
    deadline = parse_datetime(data["registration_deadline"], "Hạn đăng ký") if data.get("registration_deadline") else None
    if end <= start or (deadline and deadline > start):
        raise DomainError("Thời gian sự kiện hoặc hạn đăng ký không hợp lệ.")
    try:
        capacity = int(data["capacity"]) if data.get("capacity") is not None else None
        ban_ids = list({int(value) for value in data.get("department_ids", [])})
    except (TypeError, ValueError):
        raise DomainError("Sức chứa hoặc mã ban không hợp lệ.") from None
    if capacity is not None and capacity < 1:
        raise DomainError("Sức chứa phải lớn hơn 0.")
    applies_all = bool(data.get("applies_to_all", True))
    if not applies_all and (not ban_ids or Ban.query.filter(Ban.id.in_(ban_ids)).count() != len(ban_ids)):
        raise DomainError("Danh sách ban tham gia không hợp lệ.")
    manager_id = data.get("manager_id", current_user.id)
    manager = db.get_or_404(User, manager_id)
    require(is_board(manager), "Người phụ trách phải thuộc ban chủ nhiệm.")
    if current_user.chuc_vu == "TB" and (manager.id != current_user.id or applies_all or not set(ban_ids).issubset(current_user.quan_ly_ban_ids())):
        raise DomainError("Trưởng ban chỉ được tạo sự kiện trong ban mình quản lý.", 403)
    fund_id = data.get("fund_id")
    if fund_id:
        db.get_or_404(FundCollection, fund_id)
    event = Event(ten_su_kien=name, ma_su_kien=code, thoi_gian_bat_dau=start,
                  thoi_gian_ket_thuc=end, trang_thai="sap_dien_ra", tao_boi_id=current_user.id)
    db.session.add(event)
    db.session.flush()
    detail = EventDetail(event_id=event.id, loai=str(data.get("type", "other"))[:100],
                         noi_dung=str(data.get("content", "")), dia_diem=str(data.get("location", ""))[:300],
                         anh_dai_dien=str(data.get("image", ""))[:255], han_dang_ky=deadline,
                         so_nguoi_toi_da=capacity, phu_trach_id=manager.id, fund_id=fund_id,
                         tat_ca_thanh_vien=applies_all)
    db.session.add(detail)
    for ban_id in ([] if applies_all else ban_ids):
        db.session.add(EventTargetBan(event_id=event.id, ban_id=ban_id))
    audit(current_user, "event_create", event)
    db.session.commit()
    return jsonify(event_json(event)), 201


@api_bp.get("/events/<int:event_id>")
def event_detail(event_id):
    return jsonify(event_json(db.get_or_404(Event, event_id)))


@api_bp.patch("/events/<int:event_id>")
def event_update(event_id):
    event = db.get_or_404(Event, event_id)
    require(can_manage_event(event))
    if event.trang_thai in {"da_ket_thuc", "cancelled"}:
        raise DomainError("Không thể sửa sự kiện đã kết thúc hoặc đã hủy.")
    data = body()
    detail = db.session.get(EventDetail, event.id)
    if detail is None:
        detail = EventDetail(event_id=event.id)
        db.session.add(detail)
    if "name" in data:
        name = str(data["name"]).strip()
        if not name:
            raise DomainError("Tên sự kiện không được để trống.")
        event.ten_su_kien = name
    start = parse_datetime(data["start_at"], "Giờ bắt đầu") if "start_at" in data else event.thoi_gian_bat_dau
    end = parse_datetime(data["end_at"], "Giờ kết thúc") if "end_at" in data else event.thoi_gian_ket_thuc
    deadline = parse_datetime(data["registration_deadline"], "Hạn đăng ký") if data.get("registration_deadline") else detail.han_dang_ky
    if end <= start or (deadline and deadline > start):
        raise DomainError("Thời gian sự kiện hoặc hạn đăng ký không hợp lệ.")
    event.thoi_gian_bat_dau, event.thoi_gian_ket_thuc = start, end
    if "registration_deadline" in data:
        detail.han_dang_ky = deadline
    for key, attribute, limit in (("type", "loai", 100), ("content", "noi_dung", None),
                                   ("location", "dia_diem", 300), ("image", "anh_dai_dien", 255)):
        if key in data:
            value = str(data[key] or "").strip()
            setattr(detail, attribute, value[:limit] if limit else value)
    if "capacity" in data:
        try:
            capacity = int(data["capacity"]) if data["capacity"] is not None else None
        except (TypeError, ValueError):
            raise DomainError("Sức chứa không hợp lệ.") from None
        registered = EventRegistration.query.filter_by(event_id=event.id, trang_thai="registered").count()
        if capacity is not None and capacity < max(1, registered):
            raise DomainError("Sức chứa không được ít hơn số thành viên đã đăng ký.")
        detail.so_nguoi_toi_da = capacity
    audit(current_user, "event_update", event)
    db.session.commit()
    return jsonify(event_json(event))


@api_bp.post("/events/<int:event_id>/copy")
def event_copy(event_id):
    source = db.get_or_404(Event, event_id)
    require(can_manage_event(source))
    data = body()
    code = str(data.get("code", "")).strip()
    if not code or Event.query.filter_by(ma_su_kien=code).first():
        raise DomainError("Mã sự kiện mới không hợp lệ hoặc đã tồn tại.")
    start = parse_datetime(data.get("start_at"), "Giờ bắt đầu")
    end = parse_datetime(data.get("end_at"), "Giờ kết thúc")
    if end <= start:
        raise DomainError("Thời gian sự kiện không hợp lệ.")
    copy = Event(ten_su_kien=source.ten_su_kien, ma_su_kien=code,
                 thoi_gian_bat_dau=start, thoi_gian_ket_thuc=end,
                 trang_thai="sap_dien_ra", tao_boi_id=current_user.id)
    db.session.add(copy)
    db.session.flush()
    if source.detail:
        detail = source.detail
        db.session.add(EventDetail(event_id=copy.id, loai=detail.loai, noi_dung=detail.noi_dung,
                                   dia_diem=detail.dia_diem, anh_dai_dien=detail.anh_dai_dien,
                                   so_nguoi_toi_da=detail.so_nguoi_toi_da,
                                   phu_trach_id=current_user.id, fund_id=detail.fund_id,
                                   tat_ca_thanh_vien=detail.tat_ca_thanh_vien))
        for target in EventTargetBan.query.filter_by(event_id=source.id):
            db.session.add(EventTargetBan(event_id=copy.id, ban_id=target.ban_id))
    audit(current_user, "event_copy", copy, {"source_id": source.id})
    db.session.commit()
    return jsonify(event_json(copy)), 201


@api_bp.post("/events/<int:event_id>/cancel")
def event_cancel(event_id):
    event = db.get_or_404(Event, event_id)
    require(can_manage_event(event))
    if event.trang_thai in {"da_ket_thuc", "cancelled"}:
        raise DomainError("Sự kiện đã kết thúc hoặc đã hủy.")
    if Attendance.query.join(EventRegistration).filter(EventRegistration.event_id == event.id,
                                                         Attendance.trang_thai != "pending").first():
        raise DomainError("Không thể hủy sự kiện đã điểm danh.")
    event.trang_thai = "cancelled"
    audit(current_user, "event_cancel", event)
    db.session.commit()
    return jsonify(event_json(event))


@api_bp.post("/events/<int:event_id>/registrations")
def event_register(event_id):
    event = db.get_or_404(Event, event_id)
    data = body()
    member = db.get_or_404(User, data["member_id"]) if data.get("member_id") else current_user
    registration, created = register(current_user, event, member)
    return jsonify(id=registration.id, status=registration.trang_thai), 201 if created else 200


@api_bp.post("/registrations/<int:registration_id>/cancel")
def registration_cancel(registration_id):
    registration = db.get_or_404(EventRegistration, registration_id)
    promoted = cancel_registration(current_user, registration)
    return jsonify(status="cancelled", promoted_id=promoted.id if promoted else None)


@api_bp.get("/events/<int:event_id>/registrations")
def event_registrations(event_id):
    event = db.get_or_404(Event, event_id)
    require(is_board(current_user) or (event.detail and event.detail.phu_trach_id == current_user.id))
    query = EventRegistration.query.filter_by(event_id=event_id).order_by(EventRegistration.id)
    return listing(paginate(query), lambda row: {"id": row.id, "member_id": row.user_id,
                    "status": row.trang_thai, "attendance": row.attendance.trang_thai if row.attendance else "pending"})


@api_bp.get("/events/<int:event_id>/qr")
def event_qr(event_id):
    event = db.get_or_404(Event, event_id)
    return jsonify(token=qr_token(current_user, event), expires_in=900)


@api_bp.get("/events/<int:event_id>/qr.png")
def event_qr_png(event_id):
    event = db.get_or_404(Event, event_id)
    token = qr_token(current_user, event)
    stream = BytesIO()
    qrcode.make(token).save(stream, format="PNG")
    stream.seek(0)
    return send_file(stream, mimetype="image/png", download_name=f"event-{event.id}-checkin.png")


@api_bp.post("/attendance/checkin")
def event_checkin():
    attendance = qr_checkin(current_user, body().get("token"))
    return jsonify(id=attendance.id, status=attendance.trang_thai,
                   checked_in_at=attendance.checkin_luc.isoformat())


@api_bp.put("/registrations/<int:registration_id>/attendance")
def attendance_manual(registration_id):
    registration = db.get_or_404(EventRegistration, registration_id)
    data = body()
    attendance = manual_attendance(current_user, registration, data.get("status"), data.get("note", ""))
    return jsonify(id=attendance.id, status=attendance.trang_thai)


@api_bp.post("/events/<int:event_id>/feedback")
def event_feedback(event_id):
    event = db.get_or_404(Event, event_id)
    data = body()
    row = feedback(current_user, event, data.get("rating"), data.get("comment", ""))
    return jsonify(id=row.id, rating=row.danh_gia), 201


@api_bp.post("/assistant")
def assistant():
    question = str(body().get("question", "")).strip()
    if not question or len(question) > 2000:
        raise DomainError("Câu hỏi không hợp lệ.")
    answer = get_provider().answer(current_user, question)
    audit(current_user, "assistant_query", current_user, {"intent": answer["intent"]})
    db.session.commit()
    return jsonify(answer)


def proposal_json(proposal):
    return {"id": proposal.id, "kind": proposal.kind, "status": proposal.status,
            "provider": proposal.provider, "preview": proposal_preview(proposal),
            "payload": proposal.payload, "result_ids": proposal.result_ids,
            "created_by_id": proposal.created_by_id,
            "reviewed_by_id": proposal.reviewed_by_id}


@api_bp.get("/assistant/proposals")
def assistant_proposals():
    require(is_board(current_user))
    query = AIProposal.query.order_by(AIProposal.id.desc())
    if current_user.chuc_vu == "TB":
        query = query.filter_by(created_by_id=current_user.id)
    if request.args.get("status"):
        query = query.filter_by(status=request.args["status"])
    return listing(paginate(query), proposal_json)


@api_bp.post("/assistant/proposals")
def assistant_proposal_create():
    proposal = create_proposal(current_user, body())
    return jsonify(proposal_json(proposal)), 201


@api_bp.post("/assistant/proposals/<int:proposal_id>/review")
def assistant_proposal_review(proposal_id):
    proposal = db.get_or_404(AIProposal, proposal_id)
    data = body()
    review_proposal(current_user, proposal, data.get("decision"))
    return jsonify(proposal_json(proposal))


@api_bp.get("/settings/activity")
def activity_settings():
    require(is_admin(current_user))
    return jsonify(points=(db.session.get(ClubSetting, "activity_point_rules") or ClubSetting(value=DEFAULT_POINTS)).value,
                   thresholds=(db.session.get(ClubSetting, "activity_thresholds") or ClubSetting(value=DEFAULT_THRESHOLDS)).value)


@api_bp.put("/settings/activity")
def activity_settings_update():
    require(is_admin(current_user))
    data = body()
    for field, key, defaults in (("points", "activity_point_rules", DEFAULT_POINTS),
                                 ("thresholds", "activity_thresholds", DEFAULT_THRESHOLDS)):
        if field not in data:
            continue
        values = data[field]
        if not isinstance(values, dict) or set(values) != set(defaults):
            raise DomainError(f"Cấu hình {field} không hợp lệ.")
        try:
            values = {name: int(value) for name, value in values.items()}
        except (TypeError, ValueError):
            raise DomainError(f"Cấu hình {field} phải là số nguyên.") from None
        if field == "thresholds" and not (0 <= values["normal"] <= values["active"] <= 100 and values["inactive_days"] > 0):
            raise DomainError("Ngưỡng hoạt động không hợp lệ.")
        row = db.session.get(ClubSetting, key)
        if row is None:
            row = ClubSetting(key=key)
            db.session.add(row)
        row.value = values
        audit(current_user, "setting_update", row, {"key": key, "values": values})
    db.session.commit()
    return activity_settings()


@api_bp.get("/dashboard")
def dashboard():
    members = visible_members(current_user).all()
    stats = [activity_stats(member) for member in members]
    open_funds = FundCollection.query.filter_by(trang_thai="open").all()
    balances = []
    for fund in open_funds:
        for member in members:
            try:
                balances.append((member, fund, balance(fund, member)))
            except DomainError:
                pass
    required = sum((item[2]["required"] for item in balances), Decimal("0.00"))
    paid = sum((item[2]["paid"] for item in balances), Decimal("0.00"))
    remaining = sum((item[2]["remaining"] for item in balances), Decimal("0.00"))
    upcoming = Event.query.filter(Event.thoi_gian_bat_dau >= datetime.utcnow()).order_by(Event.thoi_gian_bat_dau).limit(5).all()
    return jsonify(members=len(members), active_members=sum(member_record(m).tinh_trang == "active" for m in members),
                   low_activity=sum(s["classification"] in {"Ít hoạt động", "Không hoạt động"} for s in stats),
                   fund_required=str(required), fund_paid=str(paid), fund_remaining=str(remaining),
                   fund_completion=round(float(paid / required * 100), 1) if required else 0,
                   average_participation=round(sum(s["rate"] for s in stats) / len(stats), 1) if stats else 0,
                   overdue=[{"member_id": m.id, "fund_id": f.id, "remaining": str(b["remaining"])}
                            for m, f, b in balances if b["status"] in {"unpaid", "overdue"}][:20],
                   upcoming_events=[event_json(event) for event in upcoming])


@api_bp.get("/reports/funds.xlsx")
def fund_report():
    require(is_treasurer(current_user))
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Tinh trang dong quy"
    sheet.append(["Mã thành viên", "Họ tên", "Khoản thu", "Phải đóng", "Đã đóng", "Còn thiếu", "Trạng thái"])
    for fund in FundCollection.query.order_by(FundCollection.id):
        for member in target_members(fund):
            item = balance(fund, member)
            sheet.append([member.mssv, member.ho_ten, fund.ten_khoan_thu,
                          item["required"], item["paid"], item["remaining"], item["status"]])
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name="bao_cao_quy.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@api_bp.get("/reports/members.xlsx")
def member_report():
    require(is_board(current_user) or can_view_all_members(current_user))
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Thanh vien"
    sheet.append(["Mã", "Họ tên", "Email", "Số điện thoại", "Chức vụ", "Trạng thái", "Ngày gia nhập"])
    for member in visible_members(current_user).order_by(User.id):
        info = member_json(member)
        sheet.append([info["code"], info["name"], info["email"], info["phone"], info["role"], info["status"], info["joined_at"]])
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name="danh_sach_thanh_vien.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
